"""Official Census occupation-code to stable SOC major-group mapping."""

from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from pathlib import Path


SOC_MAJOR_GROUPS = {
    "11": "Management Occupations",
    "13": "Business and Financial Operations Occupations",
    "15": "Computer and Mathematical Occupations",
    "17": "Architecture and Engineering Occupations",
    "19": "Life, Physical, and Social Science Occupations",
    "21": "Community and Social Service Occupations",
    "23": "Legal Occupations",
    "25": "Education, Training, and Library Occupations",
    "27": "Arts, Design, Entertainment, Sports, and Media Occupations",
    "29": "Healthcare Practitioners and Technical Occupations",
    "31": "Healthcare Support Occupations",
    "33": "Protective Service Occupations",
    "35": "Food Preparation and Serving Related Occupations",
    "37": "Building and Grounds Cleaning and Maintenance Occupations",
    "39": "Personal Care and Service Occupations",
    "41": "Sales and Related Occupations",
    "43": "Office and Administrative Support Occupations",
    "45": "Farming, Fishing, and Forestry Occupations",
    "47": "Construction and Extraction Occupations",
    "49": "Installation, Maintenance, and Repair Occupations",
    "51": "Production Occupations",
    "53": "Transportation and Material Moving Occupations",
    "55": "Military Specific Occupations",
}

UNMAPPED_SPECIAL_CODES = {"9830", "9920"}
CROSSWALK_COLUMNS = (
    "source_code_version", "source_code", "source_title", "source_soc_code",
    "target_code_candidates", "target_soc_major_candidates", "occupation_group",
    "occupation_group_title", "civilian_group", "mapping_rule", "mapping_status",
)


def normalize_census_code(value: object) -> str | None:
    """Normalize an Excel string/numeric code without changing its meaning."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return f"{value:04d}" if 0 <= value <= 9999 else None
    if isinstance(value, float) and value.is_integer():
        integer = int(value)
        return f"{integer:04d}" if 0 <= integer <= 9999 else None
    if isinstance(value, str) and re.fullmatch(r"\d{4}", value.strip()):
        return value.strip()
    return None


def soc_major_group(soc_code: object) -> str | None:
    """Return the SOC major-group prefix, including valid SOC X-ending codes."""
    if not isinstance(soc_code, str):
        return None
    match = re.fullmatch(r"(\d{2})-[0-9X]{4}", soc_code.strip().upper())
    return match.group(1) if match else None


def cps_occupation_variable(year: int) -> str:
    """The fixed-width field name changed in 2021, after the 2020 code break."""
    if not 2016 <= year <= 2026:
        raise ValueError(f"CPS layout year outside checked range: {year}")
    return "PEIO1OCD" if year <= 2020 else "PTIO1OCD"


def cps_occupation_code_version(year: int) -> str:
    """CPS switched from 2010 to 2018 Census OCC codes in January 2020."""
    if not 2016 <= year <= 2026:
        raise ValueError(f"CPS layout year outside checked range: {year}")
    return "census_2010" if year <= 2019 else "census_2018"


def _mapping_record(version: str, code: str, title: str, soc_code: str | None) -> dict[str, str]:
    group_number = soc_major_group(soc_code)
    if code in UNMAPPED_SPECIAL_CODES and group_number is None:
        mapping_status = "special_unmapped_no_civilian_occupation"
    elif group_number is None:
        raise ValueError(f"No official SOC major group for {version} occupation code {code}: {soc_code!r}")
    elif group_number not in SOC_MAJOR_GROUPS:
        raise ValueError(f"Unknown SOC major group {group_number} for {version} code {code}")
    else:
        mapping_status = "mapped"

    return {
        "source_code_version": version,
        "source_code": code,
        "source_title": title.strip(),
        "source_soc_code": soc_code or "",
        "target_code_candidates": code,
        "target_soc_major_candidates": group_number or "",
        "occupation_group": f"{group_number}-0000" if group_number else "",
        "occupation_group_title": SOC_MAJOR_GROUPS[group_number] if group_number else "",
        "civilian_group": str(bool(group_number and group_number != "55")).lower(),
        "mapping_rule": "official_equivalent_SOC_major_group_prefix",
        "mapping_status": mapping_status,
    }


def resolve_target_major_groups(target_codes: set[str], code_to_major: dict[str, str]) -> tuple[str | None, str]:
    """Map only when every official destination code implies one SOC major group."""
    missing = sorted(code for code in target_codes if code not in code_to_major)
    groups = sorted({code_to_major[code] for code in target_codes if code in code_to_major})
    if missing:
        return None, "unmapped_target_code_not_in_2018_list"
    if len(groups) > 1:
        return None, "unmapped_ambiguous_cross_year_major_group"
    if len(groups) == 1:
        return groups[0], "mapped"
    return None, "unmapped_no_verified_2018_bridge"


def _historical_bridge_candidates(
    workbook_path: Path,
    conversion_rates_path: Path,
    direct_matches_path: Path,
) -> tuple[dict[str, set[str]], dict[str, str], dict[str, str], dict[str, str], dict[str, tuple[str, str]]]:
    """Collect possible 2018 destinations; conversion probabilities are ignored."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - runtime-specific message
        raise RuntimeError("Building the official occupation crosswalk requires openpyxl") from exc

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    rates = load_workbook(conversion_rates_path, read_only=True, data_only=True)
    direct = load_workbook(direct_matches_path, read_only=True, data_only=True)
    if "2018 Census Occ Code List" not in workbook.sheetnames:
        raise ValueError("Official occupation workbook is missing its 2018 code list")
    code_to_major: dict[str, str] = {}
    code_to_title: dict[str, str] = {}
    code_to_soc: dict[str, str] = {}
    for row in workbook["2018 Census Occ Code List"].iter_rows(values_only=True):
        code = normalize_census_code(row[2] if len(row) > 2 else None)
        if code:
            soc = str(row[3] or "").strip()
            major = soc_major_group(soc)
            if major:
                code_to_major[code] = major
            code_to_title[code] = str(row[1] or "").strip()
            code_to_soc[code] = soc

    old_codes: dict[str, tuple[str, str]] = {}
    candidates: dict[str, set[str]] = defaultdict(set)
    for row in workbook["2010 to 2018 Crosswalk "].iter_rows(values_only=True):
        old_code = normalize_census_code(row[1] if len(row) > 1 else None)
        new_code = normalize_census_code(row[4] if len(row) > 4 else None)
        if old_code:
            old_soc = str(row[0] or "").strip()
            old_title = str(row[2] or "").strip()
            old_codes[old_code] = (old_soc, old_title)
        if old_code and new_code:
            candidates[old_code].add(new_code)

    # The official E1 table enumerates possible destinations for changed codes.
    # Use the destination code set only; never use its rate columns to split rows.
    previous_old_code: str | None = None
    for row in rates["E1 Total"].iter_rows(min_row=4, values_only=True):
        old_code = normalize_census_code(row[0] if len(row) > 0 else None)
        new_code = normalize_census_code(row[2] if len(row) > 2 else None)
        if old_code:
            previous_old_code = old_code
        if previous_old_code and new_code:
            candidates[previous_old_code].add(new_code)

    for row in direct["Table F1 Direct Match Occ"].iter_rows(min_row=4, values_only=True):
        old_code = normalize_census_code(row[0] if len(row) > 0 else None)
        new_code = normalize_census_code(row[2] if len(row) > 2 else None)
        if old_code and new_code:
            candidates[old_code].add(new_code)

    # The official code-change sheet lists all targets for a split across rows.
    previous_old_code = None
    for row in workbook["Occ Code Changes"].iter_rows(min_row=4, values_only=True):
        old_code = normalize_census_code(row[0] if len(row) > 0 else None)
        new_code = normalize_census_code(row[2] if len(row) > 2 else None)
        if old_code:
            previous_old_code = old_code
        if previous_old_code and new_code:
            candidates[previous_old_code].add(new_code)

    workbook.close()
    rates.close()
    direct.close()

    return candidates, code_to_major, code_to_title, code_to_soc, old_codes


def extract_official_crosswalk_rows(
    workbook_path: Path,
    conversion_rates_path: Path,
    direct_matches_path: Path,
) -> list[dict[str, str]]:
    """Map old codes only when their official 2018 destination major is unique.

    Destination-code sets come from Census Tables E1/F1 and the workbook's
    code-change list. E1 conversion rates are never used to allocate records.
    """
    candidates, code_to_major, code_to_title, code_to_soc, old_codes = _historical_bridge_candidates(
        workbook_path, conversion_rates_path, direct_matches_path,
    )
    rows = [
        _mapping_record("census_2018", code, title,
                        soc_major_group(code_to_soc[code]) and code_to_soc[code] or None)
        for code, title in code_to_title.items()
    ]
    for code, (old_soc, title) in old_codes.items():
        candidate_codes = candidates.get(code, set())
        group, status = resolve_target_major_groups(candidate_codes, code_to_major)
        if code in UNMAPPED_SPECIAL_CODES and candidate_codes.issubset(UNMAPPED_SPECIAL_CODES):
            status = "special_unmapped_no_civilian_occupation"
        rows.append({
            "source_code_version": "census_2010", "source_code": code,
            "source_title": title, "source_soc_code": old_soc,
            "target_code_candidates": ";".join(sorted(candidate_codes)),
            "target_soc_major_candidates": ";".join(sorted({code_to_major[c] for c in candidate_codes if c in code_to_major})),
            "occupation_group": f"{group}-0000" if group else "",
            "occupation_group_title": SOC_MAJOR_GROUPS[group] if group else "",
            "civilian_group": str(bool(group and group != "55")).lower(),
            "mapping_rule": "official_2010_to_2018_bridge_unique_SOC_major_group" if status == "mapped"
                            else "official_bridge_no_probability_allocation",
            "mapping_status": status,
        })

    seen: set[tuple[str, str]] = set()
    for item in rows:
        key = (item["source_code_version"], item["source_code"])
        if key in seen:
            raise ValueError(f"Duplicate source occupation code in official list: {key}")
        seen.add(key)

    expected_counts = {"census_2010": 540, "census_2018": 570}
    for version, expected in expected_counts.items():
        version_rows = [row for row in rows if row["source_code_version"] == version]
        if len(version_rows) != expected:
            raise ValueError(f"Expected {expected} {version} codes from official list; found {len(version_rows)}")
    old_status_counts = Counter(row["mapping_status"] for row in rows if row["source_code_version"] == "census_2010")
    expected_old_status_counts = {
        "mapped": 504,
        "unmapped_ambiguous_cross_year_major_group": 8,
        "unmapped_no_verified_2018_bridge": 26,
        "special_unmapped_no_civilian_occupation": 2,
    }
    if old_status_counts != Counter(expected_old_status_counts):
        raise ValueError(f"Unexpected official 2010-to-2018 bridge coverage: {dict(old_status_counts)}")
    return sorted(rows, key=lambda row: (row["source_code_version"], row["source_code"]))


def write_crosswalk(rows: list[dict[str, str]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=CROSSWALK_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def build_crosswalk(workbook_path: Path, output_path: Path, *, conversion_rates_path: Path,
                    direct_matches_path: Path) -> dict[str, object]:
    rows = extract_official_crosswalk_rows(workbook_path, conversion_rates_path, direct_matches_path)
    write_crosswalk(rows, output_path)
    summary: dict[str, object] = {
        "output_path": str(output_path),
        "total_codes": len(rows),
        "sources": {},
    }
    for version in ("census_2010", "census_2018"):
        version_rows = [row for row in rows if row["source_code_version"] == version]
        summary["sources"][version] = {
            "codes": len(version_rows),
            "mapped_to_any_soc_major_group": sum(row["mapping_status"] == "mapped" for row in version_rows),
            "mapping_status_counts": dict(Counter(row["mapping_status"] for row in version_rows)),
            "civilian_soc_major_groups": sorted({row["occupation_group"] for row in version_rows
                                                   if row["civilian_group"] == "true"}),
            "unmapped_codes": sorted(row["source_code"] for row in version_rows
                                      if row["mapping_status"] != "mapped"),
        }
    summary["unique_mappings_changing_major_group"] = [
        {"source_code": row["source_code"], "source_title": row["source_title"],
         "source_soc_major": row["source_soc_code"][:2],
         "target_soc_major": row["target_soc_major_candidates"],
         "occupation_group": row["occupation_group"]}
        for row in rows if row["source_code_version"] == "census_2010" and row["mapping_status"] == "mapped"
        and row["source_soc_code"][:2] != row["target_soc_major_candidates"]
    ]
    return summary
